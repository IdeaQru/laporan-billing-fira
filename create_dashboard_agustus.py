import os
import sqlite3
from copy import copy
from typing import NamedTuple, List, Any
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

ROOT = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(ROOT, 'data', 'wifi_billing.db')
TEMPLATE_PATH = os.path.join(ROOT, 'dashboard_backup.xlsx')
OUTPUT_PATH = os.path.join(ROOT, 'dashboard_agustus_injected.xlsx')


class CustomerRecord(NamedTuple):
    no: int
    area: str
    customer_id: str
    customer_name: str
    kode_paket: int
    paket_name: str
    harga: float
    cash: float
    bca: float
    bri: float
    mandiri: float
    bni: float
    status: str
    tunggakan_rp: float
    tunggakan_bulan: str


def copy_cell_style(src_cell, target_cell):
    """Copies styling (font, border, fill, number_format, alignment) from src_cell to target_cell."""
    if src_cell.has_style:
        target_cell.font = copy(src_cell.font)
        target_cell.border = copy(src_cell.border)
        target_cell.fill = copy(src_cell.fill)
        target_cell.number_format = src_cell.number_format
        target_cell.protection = copy(src_cell.protection)
        target_cell.alignment = copy(src_cell.alignment)


def get_agustus_records() -> List[CustomerRecord]:
    """Queries August 2026 customer billing records from SQLite."""
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute('''
        SELECT 
            c.customer_code, 
            c.name, 
            a.name, 
            p.code, 
            i.amount, 
            i.status, 
            i.unpaid_amount, 
            i.notes, 
            pm.payment_method, 
            pm.amount_paid
        FROM customers c
        JOIN areas a ON a.id = c.area_id
        LEFT JOIN packages p ON p.id = c.package_id
        JOIN invoices i ON i.customer_id = c.id AND i.billing_period = '2026-08'
        LEFT JOIN payments pm ON pm.invoice_id = i.id
        ORDER BY a.id, c.id
    ''')
    rows = cur.fetchall()
    conn.close()

    records = []
    for idx, r in enumerate(rows):
        cid, name, area, pkg_code, tariff, status, unpaid, notes, pay_method, amount_paid = r

        # Map tariff to appropriate package code if not default
        t = int(tariff)
        if t == 50000:
            k_code = 5
            p_name = '5mbps'
        elif t == 75000:
            k_code = 7
            p_name = '7mbps'
        elif t == 90000:
            k_code = 9
            p_name = '9mbps'
        elif t == 150000:
            k_code = 15
            p_name = '15mbps'
        elif t == 200000:
            k_code = 20
            p_name = '20mbps'
        elif t == 250000:
            k_code = 30
            p_name = '30mbps'
        else:
            k_code = 10
            p_name = '10mbps'

        paid = float(amount_paid or 0)
        cash = paid if pay_method == 'CASH' else 0.0
        bca  = paid if pay_method == 'BCA' else 0.0
        bri  = paid if pay_method == 'BRI' else 0.0
        mandiri = paid if pay_method == 'MANDIRI' else 0.0
        bni  = paid if pay_method == 'BNI' else 0.0

        is_paid = status == 'LUNAS'
        is_free = status == 'FREE'
        tung_rp = float(unpaid or 0) if not is_paid and not is_free else 0.0
        tung_bln = '1' if tung_rp > 0 else ''

        records.append(CustomerRecord(
            no=idx + 1,
            area=area,
            customer_id=cid,
            customer_name=name,
            kode_paket=k_code,
            paket_name=p_name,
            harga=float(tariff),
            cash=cash,
            bca=bca,
            bri=bri,
            mandiri=mandiri,
            bni=bni,
            status=status,
            tunggakan_rp=tung_rp,
            tunggakan_bulan=tung_bln
        ))

    return records


def update_setup_sheet(ws_setup: Worksheet):
    """Ensures packages for 50k, 75k, 90k, 150k exist in Setup sheet for VLOOKUP resolution."""
    packages = [
        (5, '5mbps', 50000),
        (7, '7mbps', 75000),
        (9, '9mbps', 90000),
        (10, '10mbps', 100000),
        (15, '15mbps', 150000),
        (20, '20mbps', 200000),
        (30, '30mbps', 250000),
        (50, '50mbps', 350000),
        (100, '100mbps', 1850000),
        (200, '200Mbps', 3700000),
        (300, '300Mbps', 5550000),
        (400, '400Mbps', 7400000),
        (500, '500Mbps', 9250000),
        (600, '600Mbps', 11100000),
    ]
    ref_row = 3
    ref_styles = [ws_setup.cell(ref_row, c) for c in range(1, 6)]

    for idx, (code, speed, price) in enumerate(packages):
        row_idx = 3 + idx
        for c in range(1, 6):
            copy_cell_style(ref_styles[c-1], ws_setup.cell(row_idx, c))
        ws_setup.cell(row_idx, 2).value = code
        ws_setup.cell(row_idx, 3).value = speed
        ws_setup.cell(row_idx, 4).value = price


def build_agustus_dashboard():
    print(f"[INFO] Membaca data Agustus 2026 dari: {DB_PATH}")
    records = get_agustus_records()
    print(f"[INFO] Total pelanggan Agustus: {len(records)}")

    print(f"[INFO] Membuka template: {TEMPLATE_PATH}")
    wb = openpyxl.load_workbook(TEMPLATE_PATH)

    # 1. Update Setup Sheet
    if 'Setup' in wb.sheetnames:
        update_setup_sheet(wb['Setup'])
        print("[INFO] Sheet 'Setup' diperbarui dengan konfigurasi tarif lengkap (50k, 75k, 90k, 100k, dst.)")

    # 2. Injeksi Sheet NamaPelanggan
    ws_np: Worksheet = wb['NamaPelanggan']
    np_template_row = 3
    np_styles = [ws_np.cell(np_template_row, c) for c in range(1, 9)]

    np_start_row = 3
    for idx, rec in enumerate(records):
        curr_row = np_start_row + idx
        for c in range(1, 9):
            copy_cell_style(np_styles[c-1], ws_np.cell(curr_row, c))

        ws_np.cell(curr_row, 1).value = None
        ws_np.cell(curr_row, 2).value = rec.no
        ws_np.cell(curr_row, 3).value = f'=IF(D{curr_row}="","",LEFT(SUBSTITUTE(UPPER(E{curr_row})," ",""),3)&"-"&TEXT(B{curr_row},"000"))'
        ws_np.cell(curr_row, 4).value = rec.customer_name
        ws_np.cell(curr_row, 5).value = rec.area
        ws_np.cell(curr_row, 6).value = rec.kode_paket
        ws_np.cell(curr_row, 7).value = f'=IFERROR(VLOOKUP(F{curr_row},Setup!B:C,2,FALSE),"")'
        ws_np.cell(curr_row, 8).value = f'=IFERROR(VLOOKUP(G{curr_row},Setup!C:D,2,FALSE),"")'

    # Bersihkan sisa baris lama di NamaPelanggan
    max_np_row = max(ws_np.max_row, np_start_row + len(records) + 50)
    for r in range(np_start_row + len(records), max_np_row + 1):
        for c in range(1, 9):
            ws_np.cell(r, c).value = None

    print(f"[INFO] Injeksi {len(records)} pelanggan ke sheet 'NamaPelanggan' berhasil")

    # 3. Injeksi Sheet Tagihan Pelanggan
    ws_tp: Worksheet = wb['Tagihan Pelanggan']
    tp_template_row = 4
    tp_styles = [ws_tp.cell(tp_template_row, c) for c in range(1, 15)]

    tp_start_row = 4
    for idx, rec in enumerate(records):
        curr_tp_row = tp_start_row + idx
        curr_np_row = np_start_row + idx

        for c in range(1, 15):
            copy_cell_style(tp_styles[c-1], ws_tp.cell(curr_tp_row, c))

        ws_tp.cell(curr_tp_row, 1).value = None
        ws_tp.cell(curr_tp_row, 2).value = rec.no
        ws_tp.cell(curr_tp_row, 3).value = f'=NamaPelanggan!C{curr_np_row}'
        ws_tp.cell(curr_tp_row, 4).value = f'=NamaPelanggan!D{curr_np_row}'
        ws_tp.cell(curr_tp_row, 5).value = f'=NamaPelanggan!E{curr_np_row}'
        ws_tp.cell(curr_tp_row, 6).value = f'=NamaPelanggan!H{curr_np_row}'

        ws_tp.cell(curr_tp_row, 7).value = rec.cash if rec.cash > 0 else None
        ws_tp.cell(curr_tp_row, 8).value = rec.bca if rec.bca > 0 else None
        ws_tp.cell(curr_tp_row, 9).value = rec.bri if rec.bri > 0 else None
        ws_tp.cell(curr_tp_row, 10).value = rec.mandiri if rec.mandiri > 0 else None
        ws_tp.cell(curr_tp_row, 11).value = rec.bni if rec.bni > 0 else None

        ws_tp.cell(curr_tp_row, 12).value = f'=IF(SUM(G{curr_tp_row}:K{curr_tp_row})>0,"LUNAS","BELUM LUNAS")'
        ws_tp.cell(curr_tp_row, 13).value = rec.tunggakan_rp if rec.tunggakan_rp > 0 else None
        ws_tp.cell(curr_tp_row, 14).value = rec.tunggakan_bulan if rec.tunggakan_bulan else None

    # Bersihkan sisa baris lama di Tagihan Pelanggan
    max_tp_row = max(ws_tp.max_row, tp_start_row + len(records) + 50)
    for r in range(tp_start_row + len(records), max_tp_row + 1):
        for c in range(1, 15):
            ws_tp.cell(r, c).value = None

    print(f"[INFO] Injeksi {len(records)} baris tagihan ke sheet 'Tagihan Pelanggan' berhasil")

    # 4. Pengeluaran Sheet (Agustus 2026: Total Rp 1.515.000)
    if 'Pengeluaran' in wb.sheetnames:
        ws_exp: Worksheet = wb['Pengeluaran']
        agustus_expenses = [
            ('14 agustus 2026', 'fee mas fany', 335000),
            ('16 agustus 2026', 'fee mas fany', 180000),
            ('16 agustus 2026', 'fee mba ida', 575000),
            ('16 agustus 2026', 'fee mba dita', 285000),
            ('16 agustus 2026', 'fee mba dita', 35000),
            ('23 agustus 2026', 'fee mas fany', 60000),
            ('24 agustus 2026', 'fee mba yeni', 45000),
        ]
        # Pre-capture template cell styles from Row 3
        exp_styles = [ws_exp.cell(3, c) for c in range(1, 7)]

        for r in range(3, 101):
            for c in range(2, 6):
                ws_exp.cell(r, c).value = None

        for idx, (tgl, uraian, jml) in enumerate(agustus_expenses):
            curr_r = 3 + idx
            for c in range(1, 7):
                copy_cell_style(exp_styles[c-1], ws_exp.cell(curr_r, c))
            ws_exp.cell(curr_r, 2).value = idx + 1
            ws_exp.cell(curr_r, 3).value = tgl
            ws_exp.cell(curr_r, 4).value = uraian
            ws_exp.cell(curr_r, 5).value = jml
            ws_exp.cell(curr_r, 5).number_format = '[$Rp]#,##0.00'

        ws_exp.cell(3, 6).value = '=SUM(E3:E93)'
        print(f"[INFO] Sheet 'Pengeluaran' diinjeksi 7 transaksi Agustus (Total: Rp 1.515.000)")

    # 5. Dashboard Sheet (Formula Verification & Title)
    if 'Dashboard' in wb.sheetnames:
        ws_dash: Worksheet = wb['Dashboard']
        ws_dash.cell(2, 2).value = 'DASHBOARD PEMBAYARAN TAGIHAN WIFI - AGUSTUS 2026'
        last_row = tp_start_row + len(records) - 1  # 461
        ws_dash.cell(8, 3).value = f'=SUM(\'Tagihan Pelanggan\'!M4:M{last_row})'  # Belum lunas
        ws_dash.cell(8, 5).value = f'=COUNTA(NamaPelanggan!C3:C{last_row + 1})'   # Total pelanggan
        ws_dash.cell(8, 9).value = '=Pengeluaran!F3'                            # Pengeluaran
        ws_dash.cell(8, 12).value = f'=SUM(\'Tagihan Pelanggan\'!G4:G{last_row})' # CASH
        ws_dash.cell(9, 12).value = f'=SUM(\'Tagihan Pelanggan\'!H4:H{last_row})' # BCA
        ws_dash.cell(10, 12).value = f'=SUM(\'Tagihan Pelanggan\'!I4:I{last_row})' # BRI
        ws_dash.cell(11, 12).value = f'=SUM(\'Tagihan Pelanggan\'!J4:J{last_row})' # MANDIRI
        ws_dash.cell(12, 12).value = f'=SUM(\'Tagihan Pelanggan\'!K4:K{last_row})' # BNI
        ws_dash.cell(15, 12).value = '=L8+L9+L10+L11+L12'                         # TOTAL KESELURUHAN LUNAS

    wb.save(OUTPUT_PATH)
    print(f"[SUCCESS] Berkas dashboard berhasil dibuat di: {OUTPUT_PATH}")


if __name__ == '__main__':
    build_agustus_dashboard()
